# -*- coding: utf-8 -*-
"""04_附件一 energyMatrix V1.0 典型面积配置方案 xlsx"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT = r"D:\work\yiliaohouqin\haystack code\energy\docs\releasedoc\04_附件一_energyMatrix V1.0典型面积配置方案_20260914.xlsx"

wb = Workbook()
ws = wb.active
ws.title = "典型面积配置"
ws.sheet_view.showGridLines = False

TEAL = "0B5F77"; LIGHT = "F2F7F9"; INK = "1E2832"
thin = Side(style="thin", color="D5DDE3")
bottom = Border(bottom=thin)

ws.merge_cells("B2:I2")
ws["B2"] = "energyMatrix 建筑能源管理系统 V1.0 典型面积配置方案"
ws["B2"].font = Font(name="微软雅黑", size=15, bold=True, color=INK)
ws.row_dimensions[2].height = 30
ws.merge_cells("B3:I3")
ws["B3"] = "按建筑面积分 300 / 1,000 / 2,000 ㎡ 三档；数量为典型估算，实际以项目点位清单与信息收集表（附件五）为准。西门子（中国）有限公司 · 2026-09"
ws["B3"].font = Font(name="微软雅黑", size=9, color="5C6B78")
ws.row_dimensions[3].height = 16

headers = ["序号", "设备类型", "设备名称", "设备功能描述", "单位", "300平米配置", "1000平米配置", "2000平米配置"]
for j, h in enumerate(headers):
    c = ws.cell(row=5, column=2 + j, value=h)
    c.font = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=TEAL)
    c.alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[5].height = 22

rows = [
 ("边缘引擎", "F200", "边缘引擎（CFG3.F200），内置 FIN Framework 与 energyMatrix，3 网口 + 4G 无线上网，8GB eMMC，单台最大 3,000 点；免服务器部署", "台", 1, 1, 1),
 ("边缘引擎", "NEXIO", "边缘控制器，内置 FIN Framework 与 energyMatrix；四核 A72 / 4GB / 16GB eMMC，2×RS-485 + M-Bus，单台最大 5,000 点；中大型站点或多总线混合接入时选用", "台", 0, 0, 0),
 ("软件授权", "FIN Framework 点位", "配置的 FIN Framework 最大点位数量（含表计、传感器、计算点与历史点）", "点", 500, 1000, 2000),
 ("软件授权", "energyMatrix 应用", "建筑能源管理应用授权：分项计量树、能耗台账、指标定额、双碳核算、诊断闭环、13 个业务屏", "套", 1, 1, 1),
 ("电计量", "关口多功能电表", "产权分界关口计量：双向有功、分时 TOU、最大需量、谐波，RS-485 / Modbus 上报，结算级精度", "块", 1, 1, 2),
 ("电计量", "分项 / 子表（导轨式多功能电表）", "按分项导则 A/B/C/D 四类与楼层 / 租户计量：正向有功、需量，RS-485 / Modbus", "块", 4, 12, 24),
 ("电计量", "电流互感器（CT）", "配合导轨式电表，变比录入 emCtRatio 由系统自动还原读数", "只", 12, 36, 72),
 ("水计量", "智能水表", "自来水 / 中水累计流量，M-Bus 或 RS-485 接口，支持远传抄表", "块", 1, 2, 4),
 ("气计量", "智能燃气表", "燃气累计流量（标况折算），RS-485 或 LoRa 无线上报", "块", 0, 1, 1),
 ("冷热量", "超声波冷热量表", "冷量 / 热量计量：累积能量、瞬时功率、供回水温度，M-Bus / RS-485", "台", 0, 1, 2),
 ("采集通信", "Modbus 网关 / 采集器", "RS-485 总线采集转以太网，边缘引擎 Modbus TCP 轮询；每总线不超过 32 块表", "台", 1, 2, 4),
 ("采集通信", "M-Bus 采集器", "水表 / 冷热量表 M-Bus 总线采集，NEXIO 可直连 5 台以内免采集器", "台", 0, 1, 2),
 ("采集通信", "4G DTU", "无内网区域的表计数据回传（独立电井、室外泵房等）", "台", 0, 1, 1),
 ("末端用电", "智能墙壁插座（计量型）", "末端回路用电量采集与远程通断，Zigbee 3.0，经无线协转网关接入", "个", 2, 4, 8),
 ("现场交互", "嵌入式触摸屏", "能源运行监控中心本地值守：实时负荷、告警事件流、平衡校核状态", "台", 0, 1, 1),
 ("辅材", "导轨 / 端子 / 线缆辅材包", "DIN 导轨、屏蔽双绞线（RS-485 专用）、端子与标签", "套", 1, 1, 2),
]

r = 6
for i, (cat, name, desc, unit, a, b, c3) in enumerate(rows, 1):
    vals = [i, cat, name, desc, unit, a, b, c3]
    for j, v in enumerate(vals):
        cell = ws.cell(row=r, column=2 + j, value=v)
        cell.font = Font(name="微软雅黑", size=9.5, color=INK)
        cell.border = bottom
        if j == 3:
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        elif j in (0, 4, 5, 6, 7):
            cell.alignment = Alignment(horizontal="center", vertical="center")
        else:
            cell.alignment = Alignment(vertical="center")
        if r % 2 == 0:
            cell.fill = PatternFill("solid", fgColor=LIGHT)
    ws.row_dimensions[r].height = 34
    r += 1

note_r = r + 1
ws.merge_cells(start_row=note_r, start_column=2, end_row=note_r, end_column=9)
ws.cell(row=note_r, column=2, value="说明：NEXIO 与 F200 二选一；表计数量按“关口 + 分项 + 楼层/租户”三级计量树典型值估算，虚表与缺口表由系统公式生成，不占硬件。缺口率超 5% 将阻断月度关账，配置时请保证分项回路全覆盖。")
ws.cell(row=note_r, column=2).font = Font(name="微软雅黑", size=9, color="5C6B78")
ws.cell(row=note_r, column=2).alignment = Alignment(wrap_text=True, vertical="top")
ws.row_dimensions[note_r].height = 34

widths = [6, 10, 26, 58, 6, 12, 12, 12]
for j, w in enumerate(widths):
    ws.column_dimensions[get_column_letter(2 + j)].width = w

wb.save(OUT)
print("OK", OUT)
