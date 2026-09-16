import sys, os
from pathlib import Path

SRC = Path(r"D:\work\yiliaohouqin\haystack code\energy\docs\3rddoc\01_AIoT Lighting\01_AIoT Lighting\V1.0方案资料")
OUT = Path(r"D:\work\yiliaohouqin\haystack code\energy\docs\releasedoc\_extract")

def pptx_text(p):
    from pptx import Presentation
    prs = Presentation(str(p))
    lines = []
    for i, slide in enumerate(prs.slides, 1):
        lines.append(f"\n===== Slide {i} =====")
        for shape in slide.shapes:
            if shape.has_text_frame:
                t = shape.text_frame.text.strip()
                if t: lines.append(t)
            if shape.has_table:
                for row in shape.table.rows:
                    lines.append(" | ".join(c.text.strip() for c in row.cells))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            lines.append("[notes] " + slide.notes_slide.notes_text_frame.text.strip())
    return "\n".join(lines)

def docx_text(p):
    import docx
    d = docx.Document(str(p))
    lines = []
    for para in d.paragraphs:
        t = para.text.strip()
        if t:
            style = para.style.name if para.style else ""
            prefix = f"[{style}] " if "Heading" in style or "标题" in style else ""
            lines.append(prefix + t)
    for ti, table in enumerate(d.tables):
        lines.append(f"\n--- table {ti+1} ---")
        for row in table.rows:
            lines.append(" | ".join(c.text.strip().replace("\n"," ") for c in row.cells))
    return "\n".join(lines)

def pdf_text(p):
    import pdfplumber
    lines = []
    with pdfplumber.open(str(p)) as pdf:
        for i, page in enumerate(pdf.pages, 1):
            lines.append(f"\n===== Page {i} =====")
            t = page.extract_text() or ""
            lines.append(t)
    return "\n".join(lines)

def xls_text(p):
    import xlrd
    wb = xlrd.open_workbook(str(p))
    lines = []
    for sh in wb.sheets():
        lines.append(f"\n===== Sheet: {sh.name} ({sh.nrows}x{sh.ncols}) =====")
        for r in range(min(sh.nrows, 200)):
            vals = [str(sh.cell_value(r,c)).strip() for c in range(sh.ncols)]
            if any(vals):
                lines.append(" | ".join(vals))
    return "\n".join(lines)

handlers = {".pptx": pptx_text, ".docx": docx_text, ".pdf": pdf_text, ".xls": xls_text}
for f in sorted(SRC.iterdir()):
    ext = f.suffix.lower()
    if ext in handlers:
        try:
            txt = handlers[ext](f)
            out = OUT / (f.stem + ".txt")
            out.write_text(txt, encoding="utf-8")
            print(f"OK  {f.name} -> {out.name} ({len(txt)} chars)")
        except Exception as e:
            print(f"ERR {f.name}: {e}")
