# -*- coding: utf-8 -*-
import sys, copy
sys.stdout.reconfigure(encoding="utf-8")
from docx import Document
from docx.text.paragraph import Paragraph

p = r"docs/designdoc/energyMatrix-fin-pod-detailed-design.docx"
d = Document(p)

def set_cell(cell, text):
    # keep first paragraph, wipe extra paragraphs and runs
    for para in cell.paragraphs[1:]:
        para._p.getparent().remove(para._p)
    para = cell.paragraphs[0]
    for r in list(para.runs):
        r._r.getparent().remove(r._r)
    para.add_run(text)

# 1. 封面版本
assert d.paragraphs[5].text.strip() == "版本 V0.1.1"
for r in list(d.paragraphs[5].runs): r._r.getparent().remove(r._r)
d.paragraphs[5].add_run("版本 V0.1.2")

# 2. 文档信息表
t0 = d.tables[0]
set_cell(t0.rows[3].cells[1], "V0.1.2")
up = t0.rows[5].cells[1]
set_cell(up, "AI4B-EM-DS-2026-001《energyMatrix 语义模型设计说明书》V0.1.2")

# 3. 修订记录加行
t1 = d.tables[1]
new_tr = copy.deepcopy(t1.rows[3]._tr)
t1.rows[3]._tr.addnext(new_tr)
row = t1.rows[4]
for cell, txt in zip(row.cells, ["V0.1.2", "2026-09",
    "新增核心 KPI 考核屏（/kpi）设计：复用 emKpiDefs/emKpiComputeAll/emQuotaProgressAll，站点模型新增 emBeds（核定床位数）且参数 UI 可配置；版本统一为 0.1.2",
    "架构组"]):
    set_cell(cell, txt)

# 4. 4.4 契约段后插入 KPI 契约段
src = d.paragraphs[118]
assert "能流图（Sankey" in src.text
new_p = copy.deepcopy(src._p)
src._p.addnext(new_p)
np = Paragraph(new_p, src._parent)
for r in list(np.runs): r._r.getparent().remove(r._r)
np.add_run("KPI 考核（V0.1.2 新增，见 5.3）数据契约：指标定义取自 emKpiDefs()（emKpiCode、dis、emFormula、unit、emGranularity、emDimension、emHigherIsBetter、emStandardRef），数值批量取自 emKpiComputeAll(kpiCode, [subjectRef...], span)（emFormula 走 Axon 求值，分母缺失返回 null 不当 0），定额进度取自 emQuotas(subjectRef) 与 emQuotaProgressAll(subjectRef, span)。站点考核参数（area、emCoolArea、emOccupancy、emBeds）由模型配置屏经通用写接口 emEntityUpdate(entityId, changes) 维护（admin、枚举校验、留审计）。KPI 屏只读，不新增 Axon 函数、不新增后端接口。")

# 5. 5.2 路由骨架表（单格多段）cell 末尾加 /kpi 行
t21 = d.tables[21]
cell = t21.rows[0].cells[0]
assert "/flow" in cell.text
tpl = cell.paragraphs[-1]
newp = copy.deepcopy(tpl._p)
tpl._p.addnext(newp)
np2 = Paragraph(newp, tpl._parent)
for r in list(np2.runs): r._r.getparent().remove(r._r)
np2.add_run("/kpi                         KPI 考核（核心指标 + 定额进度）")

# 6. 5.3 模块表加行（插在「能流图」行后）
t22 = d.tables[22]
assert t22.rows[2].cells[0].text.strip() == "能流图"
new_tr = copy.deepcopy(t22.rows[2]._tr)
t22.rows[2]._tr.addnext(new_tr)
row = t22.rows[3]
for cell, txt in zip(row.cells, [
    "KPI 考核",
    "核心指标卡（单位面积能耗/电耗、人均能耗/水耗、单位床位能耗、碳强度）、定额进度对标、绿色医院评审数据支撑",
    "无",
    "分母缺失显示「—」并引导至模型配置屏；参数可配置（emBeds 等站点属性）"]):
    set_cell(cell, txt)

# 7. 5.4 契约段后插入 KPI 前端契约段（定位：含「能流图前端取数契约」的段）
idx = None
for i, para in enumerate(d.paragraphs):
    if para.text.startswith("能流图前端取数契约"):
        idx = i; break
assert idx is not None
src = d.paragraphs[idx]
new_p = copy.deepcopy(src._p)
src._p.addnext(new_p)
np3 = Paragraph(new_p, src._parent)
for r in list(np3.runs): r._r.getparent().remove(r._r)
np3.add_run("KPI 考核前端取数契约（emApi 薄封装，V0.1.2）：emKpiDefs() 取指标定义、emKpiComputeAll(code, [site], span) 批量取值、emQuotaProgressAll(site, span) 取定额进度；分母缺失的指标显示「—」不当 0。")

# 8. 附录 B 加路由行
t33 = d.tables[33]
assert t33.rows[-1].cells[0].text.strip() == "/em/flow"
new_tr = copy.deepcopy(t33.rows[-1]._tr)
t33.rows[-1]._tr.addnext(new_tr)
row = t33.rows[-1]
for cell, txt in zip(row.cells, ["/em/kpi", "KpiPage", "em:read", "KPI 考核：核心指标卡与定额进度，分母缺失显示「—」"]):
    set_cell(cell, txt)

d.save(p)
print("docx saved OK")