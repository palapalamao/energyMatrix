# -*- coding: utf-8 -*-
import sys, copy
sys.stdout.reconfigure(encoding="utf-8")
from docx import Document
from docx.oxml.ns import qn

PATH = "output/energyMatrix-semantic-model-design-spec(1).docx"
d = Document(PATH)

def set_para_text(p, text):
    # keep first run formatting, drop others
    runs = p.runs
    if not runs:
        p.add_run(text); return
    runs[0].text = text
    for r in runs[1:]:
        r._element.getparent().remove(r._element)

# 1. cover version
for p in d.paragraphs:
    if p.text.strip() == "版本 V0.1.1":
        set_para_text(p, "版本 V0.1.2"); break

# 2. info table
t0 = d.tables[0]
for row in t0.rows:
    k = row.cells[0].text.strip()
    if k == "版本":
        set_para_text(row.cells[1].paragraphs[0], "V0.1.2")
    elif k == "Xeto 库":
        set_para_text(row.cells[1].paragraphs[0], "em 0.1.2")

# 3. revision table: find by header
rev_tb = None
for tb in d.tables:
    if "修订内容" in tb.rows[0].cells[2].text if len(tb.columns) > 2 else False:
        pass
for tb in d.tables:
    hdr = [c.text.strip() for c in tb.rows[0].cells]
    if "修订内容" in hdr:
        rev_tb = tb; break
assert rev_tb is not None, "revision table not found"
new_tr = copy.deepcopy(rev_tb.rows[-1]._tr)
rev_tb._tbl.append(new_tr)
# refresh row objects
from docx.table import _Row
r = _Row(new_tr, rev_tb)
vals = ["V0.1.2", "2026-09", "新增核心 KPI 考核界面需求（绿色医院评审支撑）；EmSite 增加 emBeds 床位参数；版本统一为 0.1.2", "架构组"]
for cell, v in zip(r.cells, vals):
    set_para_text(cell.paragraphs[0], v)

# 4. TOC entry after 7.1 line
for p in d.paragraphs:
    if p.text.strip().startswith("7.1") and "能流图" in p.text and "\t" in p.text:
        new_p = copy.deepcopy(p._element)
        p._element.addnext(new_p)
        from docx.text.paragraph import Paragraph
        np = Paragraph(new_p, p._parent)
        set_para_text(np, "7.2  核心 KPI 考核（绿色医院评审支撑）\t19")
        break

# 5. body: insert 7.2 before Heading1 附录 A
paras = d.paragraphs
h2_proto = None   # heading 2 prototype (7.1 title)
body_proto = None # normal body paragraph prototype
for p in paras:
    t = p.text.strip()
    if p.style is not None and p.style.name == "Heading 2" and t.startswith("7.1"):
        h2_proto = p
    if t.startswith("能流图以桑基图"):
        body_proto = p
assert h2_proto is not None and body_proto is not None

anchor = None
for p in d.paragraphs:
    if p.style is not None and p.style.name == "Heading 1" and p.text.strip().startswith("附录 A"):
        anchor = p; break
assert anchor is not None

BLOCKS = [
 ("H2", "7.2  核心 KPI 考核（绿色医院评审支撑）"),
 ("B",  "面向绿色医院评审与机构绩效考核，自动计算单位能耗强度类核心指标，给出同比、环比与达标判定，为多院区并列排名提供统一口径。本需求随 V0.1.2 新增，落地为前端新屏（hash 路由 /kpi，FIN 顶栏菜单「KPI 考核」，归「分析优化」组，位于「能耗分析」与「定额与对标」之间），是第 15 屏。"),
 ("B",  "指标清单（初版，全部由既有 EmKpi 指标定义驱动，评审口径调整后可在「指标定义表」配置扩展，不改代码）："),
 ("B",  "单位建筑面积综合能耗 EUI（kWh/m²·月）——既有定义 EUI_TOTAL；单位面积电耗 EUI_ELEC（kWh/m²·月）"),
 ("B",  "单位建筑面积水耗 WUI（m³/m²·月）——既有定义 WUI"),
 ("B",  "人均能耗（kWh/人·月）——既有定义 EUI_PC；人均水耗（m³/人·月）——既有定义 WATER_PER_BED（现有编码名与中文名不符，V0.1.2 一并校正）"),
 ("B",  "单位床位能耗（kWh/床·月）——新增；绿色医院评审核心指标，依赖新模型参数 emBeds（见下）"),
 ("B",  "单位面积碳排（kgCO₂/m²·年）——既有定义 CBEI，按年粒度"),
 ("B",  "数据口径：Layer 2 能耗台账（铁律 6，禁直读 L1 点位历史）。指标值一律由后端 EmKpiService 求值（emFormula 为 Axon 表达式，分母缺失返回 null 不得当 0），前端只做展示、不自算口径。粒度以月为主，碳排强度按年。"),
 ("B",  "模型变更（本需求唯一模型缺口）：EmSite 增加 emBeds:N（在册床位数，单位床位能耗的分母），Xeto 定义与 demo 数据同步补齐（5 个站点）；既有参数 area（建筑面积）、emOccupancy（在册人数）、emCoolArea（空调面积）沿用，不重定义。"),
 ("B",  "筛选维度：账期（月，可回看任意已出账账期，支持多月趋势）；对象范围（多院区并列对比与排名；单站点下可下钻到 emDimension 允许的楼层 / 租户对象）。"),
 ("B",  "交互约定：KPI 卡片（当前值、单位、同比、环比、达标状态徽标）；院区排名表（同一指标跨对象排名，分母缺失显示「—」并注明缺哪个参数）；多月趋势图（叠加定额 / 国标基准线，基准线必须标注 emLimitSource 来源）；台账为空或指标无法求值时给空态引导，不得白屏、不得以 0 充数。"),
 ("B",  "权限与边界：本屏只读，不新增后端接口与 Axon 函数（复用既有 emKpiDefs / emKpiCompute / emKpiComputeAll / emQuotas / emQuotaProgressAll）；KPI 定义编辑与定额录入维持既有功能入口；EmKpiPoint 归一化点位固化（L3 写历史）不在 V0.1.2 范围，列后续版本。"),
]

from docx.text.paragraph import Paragraph
cur = anchor._element
first = True
for kind, text in reversed(list(enumerate(BLOCKS))):
    pass
# insert in order before anchor: keep anchor fixed, insert sequentially before it
for kind, text in BLOCKS:
    if kind == "H2":
        el = copy.deepcopy(h2_proto._element)
    else:
        el = copy.deepcopy(body_proto._element)
    cur.addprevious(el)
    np = Paragraph(el, anchor._parent)
    set_para_text(np, text)
    # blank spacer paragraph after each block (mimic doc style)
    sp = copy.deepcopy(body_proto._element)
    cur.addprevious(sp)
    Paragraph(sp, anchor._parent).runs.clear() if False else None
    from docx.oxml import OxmlElement
    # clear text of spacer
    for r in list(Paragraph(sp, anchor._parent).runs):
        r._element.getparent().remove(r._element)

d.save(PATH)
print("docx saved")