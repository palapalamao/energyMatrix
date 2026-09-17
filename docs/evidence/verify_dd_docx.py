# -*- coding: utf-8 -*-
import sys
sys.stdout.reconfigure(encoding="utf-8")
from docx import Document
d = Document(r"docs/designdoc/energyMatrix-fin-pod-detailed-design.docx")
full = "\n".join(p.text for p in d.paragraphs)
for tb in d.tables:
    for r in tb.rows:
        full += "\n" + " | ".join(c.text for c in r.cells)
print("V0.1.2 count:", full.count("V0.1.2"))
print("cover 版本 V0.1.2:", "版本 V0.1.2" in full)
print("V0.1.1 remaining (should be 4: 2 historical rows + sankey contract x2):", full.count("V0.1.1"))
print("emBeds count:", full.count("emBeds"))
print("西门子中国 count:", full.count("西门子中国"))
print("和碳 remaining:", full.count("和碳"))
print("has KPI contract 4.4:", "KPI 考核（V0.1.2 新增，见 5.3）数据契约" in full)
print("has KPI frontend contract:", "KPI 考核前端取数契约（emApi 薄封装，V0.1.2）" in full)
print("has module row:", "核心指标卡（单位面积能耗/电耗" in full)
print("has /em/kpi row:", "/em/kpi | KpiPage" in full)
print("has /kpi skeleton:", "/kpi                         KPI 考核（核心指标 + 定额进度）" in full)
print("has revision row:", "站点模型新增 emBeds（核定床位数）且参数 UI 可配置" in full)
# 修订记录行数
print("revision rows:", len(d.tables[1].rows))
print("module rows:", len(d.tables[22].rows), "route rows:", len(d.tables[33].rows))