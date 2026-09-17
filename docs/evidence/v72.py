import sys
sys.stdout.reconfigure(encoding="utf-8")
from docx import Document
d = Document("output/energyMatrix-semantic-model-design-spec(1).docx")
full = "\n".join(p.text for p in d.paragraphs)
for tb in d.tables:
    for r in tb.rows:
        full += "\n" + " | ".join(c.text for c in r.cells)
print("参数可配置 in docx:", "参数可配置：站点基础参数" in full)
print("修订行 updated:", "基础参数可在数据模型配置屏在线维护" in full)
print("V0.1.2 count:", full.count("V0.1.2"))