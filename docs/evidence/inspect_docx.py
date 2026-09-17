import sys
sys.stdout.reconfigure(encoding="utf-8")
from docx import Document
d = Document("output/energyMatrix-semantic-model-design-spec(1).docx")
for i, p in enumerate(d.paragraphs):
    t = p.text.strip()
    if t and ("V0.1" in t or "0.1.1" in t or "0.1.2" in t or "附录 A" in t or "7.1" in t or "7.2" in t):
        st = p.style.name if p.style is not None else "None"
        print(i, st, "|", t[:60])
print("---t0---")
tb = d.tables[0]
for r in tb.rows:
    print("   |", " | ".join(c.text.strip()[:26] for c in r.cells))