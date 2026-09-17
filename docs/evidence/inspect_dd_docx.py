# -*- coding: utf-8 -*-
import sys
sys.stdout.reconfigure(encoding="utf-8")
from docx import Document
d = Document(r"docs/designdoc/energyMatrix-fin-pod-detailed-design.docx")
print("paras:", len(d.paragraphs), "tables:", len(d.tables))
for i, p in enumerate(d.paragraphs):
    t = p.text.strip()
    if t and ("V0.1" in t or "V1.0" in t or "修订" in t or "4.4" in t or "5.2" in t or "5.3" in t or "附录 B" in t or "能流图" in t):
        st = p.style.name if p.style is not None else "None"
        print(i, st, "|", t[:70])
print("---tables---")
for ti, tb in enumerate(d.tables):
    r0 = " | ".join(c.text.strip()[:18] for c in tb.rows[0].cells)
    print(ti, len(tb.rows), "rows |", r0[:80])