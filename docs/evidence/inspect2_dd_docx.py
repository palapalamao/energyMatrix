# -*- coding: utf-8 -*-
import sys
sys.stdout.reconfigure(encoding="utf-8")
from docx import Document
d = Document(r"docs/designdoc/energyMatrix-fin-pod-detailed-design.docx")
for i in list(range(112,145)) + list(range(178,182)):
    if i < len(d.paragraphs):
        print(i, "|", repr(d.paragraphs[i].text[:80]))
print("=== table0 rows ===")
for r in d.tables[0].rows:
    print(" | ".join(c.text.strip()[:30] for c in r.cells))
print("=== table1 (修订记录) ===")
for r in d.tables[1].rows:
    print(" | ".join(c.text.strip()[:40] for c in r.cells))
print("=== table22 last rows ===")
for r in d.tables[22].rows[-3:]:
    print(" | ".join(c.text.strip()[:30] for c in r.cells))
print("=== table33 last rows ===")
for r in d.tables[33].rows[-3:]:
    print(" | ".join(c.text.strip()[:30] for c in r.cells))