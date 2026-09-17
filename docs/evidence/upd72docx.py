# -*- coding: utf-8 -*-
import io, sys, copy
sys.stdout.reconfigure(encoding="utf-8")
from docx import Document
from docx.text.paragraph import Paragraph

PATH = "output/energyMatrix-semantic-model-design-spec(1).docx"
d = Document(PATH)

def set_para_text(p, text):
    runs = p.runs
    if not runs:
        p.add_run(text); return
    runs[0].text = text
    for r in runs[1:]:
        r._element.getparent().remove(r._element)

# 1. append config paragraph after the 模型变更 paragraph
body_proto = None
anchor_para = None
for p in d.paragraphs:
    t = p.text.strip()
    if t.startswith("模型变更"):
        anchor_para = p
    if t.startswith("能流图以桑基图"):
        body_proto = p
assert anchor_para is not None and body_proto is not None

cfg = "参数可配置：站点基础参数（建筑面积 / 空调面积 / 在册人数 / 床位数）统一在既有「数据模型配置」屏在线维护（PropertyForm 表单经既有 emEntityUpdate 写入，admin 权限、拒绝受保护标签、写操作留审计），KPI 考核屏不内嵌参数编辑；参数缺失时 KPI 屏显示「—」并引导去模型配置补录。"
el = copy.deepcopy(body_proto._element)
anchor_para._element.addnext(el)
Paragraph(el, anchor_para._parent)
set_para_text(Paragraph(el, anchor_para._parent), cfg)
# spacer
sp = copy.deepcopy(body_proto._element)
el.addnext(sp)
for r in list(Paragraph(sp, anchor_para._parent).runs):
    r._element.getparent().remove(r._element)

# 2. revision row text update
for tb in d.tables:
    hdr = [c.text.strip() for c in tb.rows[0].cells]
    if "修订内容" in hdr:
        last = tb.rows[-1]
        if "新增核心 KPI 考核" in last.cells[2].text:
            set_para_text(last.cells[2].paragraphs[0], "新增核心 KPI 考核界面需求（绿色医院评审支撑）；EmSite 增加 emBeds 床位参数且基础参数可在数据模型配置屏在线维护；版本统一为 0.1.2")
        break

d.save(PATH)
print("docx saved")