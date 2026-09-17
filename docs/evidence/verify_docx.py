import sys
sys.stdout.reconfigure(encoding="utf-8")
from docx import Document
d = Document("output/energyMatrix-semantic-model-design-spec(1).docx")
full = "\n".join(p.text for p in d.paragraphs)
for tb in d.tables:
    for r in tb.rows:
        full += "\n" + " | ".join(c.text for c in r.cells)
print("V0.1.2 count:", full.count("V0.1.2"))
print("V0.1.1 remaining (should be 2: 7.1 historical + revision row):", full.count("V0.1.1"))
print("emBeds count:", full.count("emBeds"))
print("西门子中国 count:", full.count("西门子中国"))
print("和碳 remaining:", full.count("和碳"))
print("has 7.2 heading:", "7.2  核心 KPI 考核（绿色医院评审支撑）" in full)
print("has TOC entry:", "7.2  核心 KPI 考核（绿色医院评审支撑）\t19" in full)
print("has revision row:", "新增核心 KPI 考核界面需求（绿色医院评审支撑）" in full)
print("has em 0.1.2:", "em 0.1.2" in full)